"""Workstation XLSX renderer (Artifact Tool is unavailable on this host).

Only presentation and spreadsheet reconciliation live here. Detection, gates,
assignment and error calculation remain in the verified calculation modules.
"""
from __future__ import annotations

import math
import shutil
import subprocess
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


NAMES = ['汇总', '各场景', '速度分组', '到点明细', '误差分布', '异常明细', '统计样本', '计算说明']
STAT_KEYS = ['accuracy_samples', 'p90_abs_mm', 'p95_abs_mm', 'signed_mean_mm',
             'sample_std_mm', 'sample_variance_mm2', 'early', 'late']
HEADERS = ['范围', '有效轮次', 'P90 mm', 'P95 mm', '均值 mm', '样本标准差 mm', '样本方差 mm²', '偏早', '偏晚']


def table(ws, headers, rows, note=''):
    ws.cell(2, 1, ws.title)
    ws.cell(4, 1, note)
    for j, value in enumerate(headers, 1):
        ws.cell(5, j, value)
    for i, row in enumerate(rows, 6):
        for j, value in enumerate(row, 1):
            ws.cell(i, j, value)
    ws.freeze_panes = 'B6'
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = f'A5:{get_column_letter(len(headers))}{max(5, ws.max_row)}'
    for row in ws:
        for c in row:
            c.font = Font(name='Noto Sans CJK SC', size=10, color='243247')
            c.alignment = Alignment(vertical='center')
            if c.row >= 6 and (isinstance(c.value, (int, float)) or c.data_type=='f'):
                c.number_format = '0.00;[Red]-0.00'
    for i in range(6,ws.max_row+1):
        ws.row_dimensions[i].height=22
    ws['A2'].font = Font(name='Noto Sans CJK SC', size=16, bold=True, color='243A59')
    for c in ws[5]:
        c.fill = PatternFill('solid', fgColor='243A59')
        c.font = Font(name='Noto Sans CJK SC', size=10, bold=True, color='FFFFFF')
        c.alignment = Alignment(wrap_text=True, vertical='center',horizontal='center')
    ws.row_dimensions[5].height = 34
    for j in range(1, len(headers)+1):
        ws.column_dimensions[get_column_letter(j)].width = 21 if j > 1 else 27
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:5'


def create_draft(result, path):
    wb = Workbook()
    wb.remove(wb.active)
    sheets = {n: wb.create_sheet(n) for n in NAMES}
    rows = result['detail']
    source = {s['scene_id']: s['input'] for s in result['scenes']}
    labels = ['样本ID','场景','前轮','触地存在','双轮配对','配对歧义','门禁 1/0/未知',
              'StartDist FL','StartDist FR','触地 epoch ns','触地相对秒','触地速度 kph',
              'jerk 通道','距离通道','预测周期','真零 epoch ns','状态','带符号误差 mm',
              '绝对误差 mm','纳入精度 1/0','排除原因','质量标记','关联样本','触地类别',
              '左控制 epoch ns','右控制 epoch ns','左距离 mm','右距离 mm','插值间隔 ms',
              '完整场景名','事件序号','距离源文件','jerk 源文件']
    matrix = []
    for i, x in enumerate(rows, 6):
        matrix.append([x['sample_id'], x['scene_id'], x['wheel'], int(x['contact_available']),
                       int(x['paired_front_contacts']), int(x['pair_ambiguous']),
                       '未知' if x['gate_eligible'] is None else int(x['gate_eligible']),
                       x['startdist_fl'], x['startdist_fr'], str(x['contact_timestamp_ns']) if x['contact_timestamp_ns'] is not None else None,
                       x['contact_time_s'], x['speed_kph'], f"s_InDatCalcFAS_WhlJerk{x['wheel']}_mps3",
                       x['wheel'].lower()+'_dist_out_mm_0', x['detection_cycle_id'],
                       str(x['predicted_zero_timestamp_ns']) if x['predicted_zero_timestamp_ns'] is not None else None,
                       x['status'], x['error_mm'], f'=IF(ISNUMBER(R{i}),ABS(R{i}),"")',
                       f'=IF(AND(D{i}=1,G{i}=1,ISNUMBER(R{i}),S{i}<=500),1,0)',
                       x['exclusion_reason'], '; '.join(x['flags']), '; '.join(x['related_samples']),
                       x['feature_at_contact'], str(x['left_timestamp_ns']) if x['left_timestamp_ns'] is not None else None,
                       str(x['right_timestamp_ns']) if x['right_timestamp_ns'] is not None else None,
                       x['left_distance_mm'], x['right_distance_mm'], x['control_bracket_ms'], x['scene'], x['event'],
                       source[x['scene_id']]['timeline']['path'], source[x['scene_id']]['jerk'][x['wheel'].lower()]['path']])
    table(sheets['到点明细'], labels, matrix, 'ns 时间戳为精确文本；未知门禁、触地缺失与预测未匹配分别保留。')
    for c in ['J','P','Y','Z']:
        sheets['到点明细'].column_dimensions[c].width = 29
        for i in range(6, 6+len(rows)):
            sheets['到点明细'][f'{c}{i}'].number_format = '@'
    for c in ['D','E','F','G','H','I','T','X','AE']:
        for i in range(6,6+len(rows)):
            sheets['到点明细'][f'{c}{i}'].number_format='#,##0'
    for c in ['M','N','U','V','W','AD','AF','AG']:
        sheets['到点明细'].column_dimensions[c].width = 55
    scopes = [('all', '合并', list(range(len(rows))), result['overall']['combined'])]
    scopes += [(w, w, [i for i,x in enumerate(rows) if x['wheel']==w], result['overall'][w]) for w in ['FL','FR']]
    for lo, hi in [(0,10),(10,20),(20,30),(30,40)]:
        scopes.append((f'speed{lo}', f'{lo}–{hi} kph', [i for i,x in enumerate(rows) if x['speed_kph'] is not None and lo<=x['speed_kph'] and (x['speed_kph']<hi or hi==40 and x['speed_kph']==40)], result['speed_bins'][f'{lo}-{hi}_kph']))
    for scene in result['scenes']:
        for w in ['combined','FL','FR']:
            key = scene['scene_id']+w
            scopes.append((key, scene['scene_id']+(' 合并' if w=='combined' else ' '+w),
                           [i for i,x in enumerate(rows) if x['scene_id']==scene['scene_id'] and (w=='combined' or x['wheel']==w)], scene['metrics'][w]))
    helper, by_key, checks = [], {}, []
    pos = 6
    for key, name, indices, expected in scopes:
        start = pos
        for ix in indices:
            helper.append([name, rows[ix]['sample_id'], f'=IF(\'到点明细\'!T{ix+6}=1,\'到点明细\'!R{ix+6},"")', f'=IF(ISNUMBER(C{pos}),ABS(C{pos}),"")', f'=IF(\'到点明细\'!T{ix+6}=1,\'到点明细\'!Q{ix+6},"")'])
            pos += 1
        if not indices:
            helper.append([name, None, None, None, None]); pos += 1
        v, a = f"'统计样本'!C{start}:C{pos-1}", f"'统计样本'!D{start}:D{pos-1}"
        stats = [f'=COUNT({v})', f'=IF(COUNT({v})=0,"不可用",PERCENTILE({a},0.9))',
                 f'=IF(COUNT({v})=0,"不可用",PERCENTILE({a},0.95))',
                 f'=IF(COUNT({v})=0,"不可用",AVERAGE({v}))',
                 f'=IF(COUNT({v})<2,"不可用",STDEV({v}))', f'=IF(COUNT({v})<2,"不可用",VAR({v}))',
                 f'=COUNTIFS(\'统计样本\'!E{start}:E{pos-1},"early")', f'=COUNTIFS(\'统计样本\'!E{start}:E{pos-1},"late")']
        by_key[key] = (name, stats, expected)
        helper.append([None]*5); pos += 1
    table(sheets['统计样本'], ['统计范围','样本ID','纳入的带符号误差 mm','纳入的绝对误差 mm','纳入的状态'], helper,
          '各范围独立引用；整体合并轮次，不平均场景 P95。空字符串不计入统计。')
    def summary_table(sheet, keys, note):
        vals = [[by_key[k][0]]+by_key[k][1] for k in keys]
        table(sheets[sheet], HEADERS, vals, note)
        for i,k in enumerate(keys, 6):
            checks.append((sheet, i, by_key[k][2]))
            for col in ['B','H','I']:
                sheets[sheet][f'{col}{i}'].number_format='#,##0'
    summary_table('汇总', ['all','FL','FR'], f"{result['scene_count']} 场景；工作站现有回灌 CSV + 同场景 XCP；不作额外补偿。")
    last = max(6, 5+len(rows))
    def count(wheel, pairs):
        terms=[]
        if wheel!='combined':
            terms += [f"'到点明细'!C6:C{last}",f'"{wheel}"']
        for col,value in pairs:
            terms += [f"'到点明细'!{col}6:{col}{last}",value]
        return '=COUNTIFS('+','.join(terms)+')' if terms else f'=COUNTA(\'到点明细\'!A6:A{last})'
    flow_headers=['范围','预期轮次','门禁通过','检出（含超范围）','超 ±500 mm','精度样本','门禁不通过','触地缺失','门禁内未匹配','门禁内检出率']
    for j,v in enumerate(flow_headers,1):
        c=sheets['汇总'].cell(11,j,v);c.font=Font(name='Noto Sans CJK SC',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='243A59');c.alignment=Alignment(wrap_text=True)
    sheets['汇总'].row_dimensions[11].height=34
    for i,w in enumerate(['combined','FL','FR'],12):
        detected='='+count(w,[('G','1'),('Q','"early"')])[1:]+'+'+count(w,[('G','1'),('Q','"late"')])[1:]
        formulas=[count(w,[]),count(w,[('G','1')]),detected,
                  count(w,[('G','1'),('S','">500"')]),count(w,[('T','1')]),count(w,[('G','0')]),
                  count(w,[('D','0')]),count(w,[('G','1'),('Q','"missed"')]),f'=IF(C{i}=0,"不可用",D{i}/C{i})']
        sheets['汇总'].cell(i,1,'合并' if w=='combined' else w)
        for j,v in enumerate(formulas,2):
            sheets['汇总'].cell(i,j,v).number_format='0.0%' if j==10 else '#,##0'
    sheets['汇总'].cell(17,1,'预期轮次含对侧触地缺失占位；门禁检出率 = 门禁通过且检出 / 门禁通过。')
    sheets['汇总'].cell(19,1,'模式：bench_csv_trial_v1；DIT：本批不可用；双轮 StartDist 门禁启用。')
    sheets['汇总'].cell(21,1,'未提供正式验收阈值；自动 jerk 候选及物理时差未人工标定。')
    summary_table('各场景', [s['scene_id']+w for s in result['scenes'] for w in ['combined','FL','FR']], '每场景列出合并、FL、FR；无有效样本显示不可用。')
    summary_table('速度分组', [f'speed{lo}' for lo in [0,10,20,30]], '触地车速分组：[0,10)、[10,20)、[20,30)、[30,40] kph。')
    valid = [x['error_mm'] for x in rows if x['used_for_accuracy']]
    lo = math.floor(min(valid)/10)*10 if valid else 0
    hi = math.floor(max(valid)/10)*10+10 if valid else 10
    bins = [[b,b+10,str(b), f'=COUNTIFS(\'到点明细\'!T6:T{last},1,\'到点明细\'!R6:R{last},">="&A{i},\'到点明细\'!R6:R{last},"<"&B{i})'] for i,b in enumerate(range(lo,hi,10),6)]
    table(sheets['误差分布'], ['下界 mm','上界 mm','横轴 mm','轮次数'], bins, '10 mm 带符号误差分箱，左闭右开；负值偏早，正值偏晚。')
    chart = BarChart()
    chart.title = '到点误差分布（10 mm 分箱）'
    chart.y_axis.title, chart.x_axis.title = '轮次数', '带符号误差 mm（箱下界）'
    chart.legend = None
    chart.add_data(Reference(sheets['误差分布'], min_col=4,min_row=5,max_row=5+len(bins)),titles_from_data=True)
    # openpyxl emits a numeric category reference. Use numeric lower bounds;
    # LibreOffice discards a numeric reference to text category cells.
    chart.set_categories(Reference(sheets['误差分布'],min_col=1,min_row=6,max_row=5+len(bins)))
    chart.width, chart.height = 24, 13
    chart.style=10
    chart.varyColors=False
    chart.autoTitleDeleted=False
    chart.x_axis.delete=False
    chart.y_axis.delete=False
    chart.y_axis.scaling.min=0
    chart.y_axis.majorUnit=1
    chart.x_axis.tickLblPos='low'
    chart.x_axis.tickLblSkip=5
    chart.x_axis.title.tx.rich.bodyPr.rot=0
    chart.series[0].graphicalProperties.solidFill='386CB0'
    sheets['误差分布'].add_chart(chart,'F5')
    exception_indices = [i for i,x in enumerate(rows) if x['flags'] or x['gate_eligible'] is False or x['status']=='missed' or not x['contact_available'] or x['error_mm'] is not None and abs(x['error_mm'])>500]
    cols = ['A','B','C','D','G','Q','R','T','U','V','W']
    table(sheets['异常明细'], ['样本ID','场景','轮','触地存在','门禁','状态','带符号误差 mm','纳入精度','排除原因','质量标记','关联样本'],
          [[f'=IF(ISBLANK(\'到点明细\'!{c}{i+6}),"",\'到点明细\'!{c}{i+6})' for c in cols] for i in exception_indices],
          '漏检/未匹配、门禁不通过、超范围、触地缺失和质量标记；分类可重叠，标记不自动排除。')
    qrow = sheets['异常明细'].max_row+3
    sheets['异常明细'].cell(qrow,1,'原始数据异常（含不完整 CSV 行、距离跳变）')
    for i,q in enumerate([q for s in result['scenes'] for q in s['data_quality']],qrow+1):
        for j,v in enumerate([q['scene_id'],q['kind'],q.get('wheel'),str(q.get('timestamp_ns') or ''),q.get('source_line'),q.get('reason')],1):
            sheets['异常明细'].cell(i,j,v)
    notes = [[k,str(v)] for k,v in result['rules'].items()]
    notes += [['来源 Run ID', result['run_id']],['产品身份', result['product_identity']],
              ['核心 SHA-256', result['core']['sha256']],['来源收据 SHA-256', result['source_receipt']['sha256']],
              ['输入', '直接读取工作站已有预测距离、车速、StartDist 和同场景 XCP jerk；不要求轨迹/视觉评测报告'],
              ['模式', '已验证的 CSV/XCP 适配核心；尚未核对的两份工程师流程脚本不阻塞此模式'],
              ['DIT / 几何 / 补偿对比', '当前输入不可用；不推断、不补造'],
              ['统计', '绝对误差 P90/P95（PERCENTILE 等价于 PERCENTILE.INC）；带符号均值；样本标准差与方差 ddof=1'],
              ['刷新', '本表是该次输入快照；更换输入需重新生成独立评测'],
              ['风险及限制', '\n'.join(result['limitations'])]]
    notes += [[s['scene_id'],s['scene']] for s in result['scenes']]
    table(sheets['计算说明'], ['项目','规则或来源'], notes)
    sheets['计算说明'].column_dimensions['A'].width = 36
    sheets['计算说明'].column_dimensions['B'].width = 115
    for row in sheets['计算说明'].iter_rows(min_row=6):
        row[1].alignment = Alignment(wrap_text=True,vertical='top')
        sheets['计算说明'].row_dimensions[row[0].row].height = 42
    for ws in sheets.values():
        ws.print_area = f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
    sheets['误差分布'].print_area=f'A1:Q{max(34,sheets["误差分布"].max_row)}'
    for i in range(6,6+len(bins)):
        for c in ['A','B','D']:
            sheets['误差分布'][f'{c}{i}'].number_format='#,##0'
    for i in range(6,6+len(exception_indices)):
        for c in ['D','E','H']:
            sheets['异常明细'][f'{c}{i}'].number_format='#,##0'
    wb.save(path)
    return checks, len(bins), exception_indices


def reconcile(path, result, checks, bin_count, exceptions):
    data = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    if data.sheetnames != NAMES:
        raise ValueError('XLSX worksheets differ from required report')
    for ws in data:
        for row in ws:
            if any(c.data_type=='e' for c in row):
                raise ValueError('XLSX contains formula errors: '+ws.title)
    def close(value, expected):
        if expected is None:
            return value == '不可用'
        return isinstance(value,(int,float)) and math.isfinite(value) and abs(value-expected)<=1e-6
    for sheet, row, expected in checks:
        for j,key in enumerate(STAT_KEYS,2):
            if not close(data[sheet].cell(row,j).value,expected[key]):
                raise ValueError(f'XLSX metric mismatch: {sheet} {row} {key}')
    for i,w in enumerate(['combined','FL','FR'],12):
        m=result['overall'][w]
        expected=[m[k] for k in ['expected','gate_eligible','detected','out_of_range','accuracy_samples','gate_failed','contact_missing','unmatched_predictions_gated','detection_rate']]
        if not all(close(data['汇总'].cell(i,j).value,v) for j,v in enumerate(expected,2)):
            raise ValueError('XLSX population flow mismatch: '+w)
    for i,x in enumerate(result['detail'],6):
        ws = data['到点明细']
        if ws.cell(i,20).value != int(x['used_for_accuracy']):
            raise ValueError('XLSX detail population differs from calculation')
        if x['error_mm'] is not None and not close(ws.cell(i,18).value,x['error_mm']):
            raise ValueError('XLSX raw error differs from calculation')
        for c,key in [(10,'contact_timestamp_ns'),(16,'predicted_zero_timestamp_ns'),(25,'left_timestamp_ns'),(26,'right_timestamp_ns')]:
            if x[key] is not None and ws.cell(i,c).value != str(x[key]):
                raise ValueError('XLSX exact timestamp lost precision')
    if sum(data['误差分布'].cell(i,4).value for i in range(6,6+bin_count)) != result['overall']['combined']['accuracy_samples']:
        raise ValueError('XLSX histogram population mismatch')
    for n,ix in enumerate(exceptions,6):
        x=result['detail'][ix]
        expected=['未知' if x['gate_eligible'] is None else int(x['gate_eligible']),int(x['used_for_accuracy'])]
        if [data['异常明细'].cell(n,c).value for c in [5,8]] != expected:
            raise ValueError('XLSX exception zero/blank mismatch')
    if not formulas['误差分布']._charts:
        raise ValueError('XLSX distribution chart missing')
    chart=formulas['误差分布']._charts[0]
    if chart.x_axis.delete or chart.y_axis.delete or chart.series[0].cat is None:
        raise ValueError('XLSX distribution axis/category labels missing')
    if not any(c.data_type=='f' for row in formulas['汇总'] for c in row):
        raise ValueError('XLSX formulas missing')
    return {'state':'VERIFIED','all_metrics_reconciled':True,'all_raw_errors_reconciled':True,
            'exact_timestamps_verified':True,'exception_population_verified':True,'histogram_verified':True,
            'sheets':NAMES,'formula_errors':0,'chart_count':1}


def render(result, output):
    draft, recalculated = output/'workbook-draft', output/'workbook-recalculated'
    draft.mkdir(); recalculated.mkdir()
    source = draft/'arrival-accuracy.xlsx'
    checks,bins,exceptions = create_draft(result,source)
    completed = subprocess.run(['libreoffice','-env:UserInstallation='+str((output/'lo-profile').as_uri()),
                               '--headless','--convert-to','xlsx','--outdir',str(recalculated),str(source)],
                              capture_output=True,text=True,timeout=120,check=True)
    (output/'workbook.log').write_text(completed.stdout+'\n'+completed.stderr,encoding='utf-8')
    path = recalculated/source.name
    validation = reconcile(path,result,checks,bins,exceptions)
    shutil.copyfile(path,output/source.name)
    return validation
